"""
مكونات واجهة مشتركة تُستخدم في أكثر من شاشة:
  - SearchableTable: جدول بيانات مع شريط بحث علوي يفلتر النتائج فوراً،
    وإمكانية تصدير المحتوى الظاهر إلى CSV.
  - KPICard: بطاقة إحصائية صغيرة تُستخدم في لوحة التحكم.
  - ScreenHeader: عنوان موحد لكل شاشة مع وصف مختصر وزر تحديث.
"""

from __future__ import annotations

import csv
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from typing import Callable, List, Sequence

import customtkinter as ctk

from app.ui import theme


def _style_treeview():
    """يطبق مظهراً داكناً على ttk.Treeview (لا تدعمه CustomTkinter أصلاً)."""
    style = ttk.Style()
    try:
        style.theme_use("clam")
    except Exception:
        pass

    style.configure(
        "Dark.Treeview",
        background=theme.COLOR_BG_PANEL,
        fieldbackground=theme.COLOR_BG_PANEL,
        foreground=theme.COLOR_TEXT_PRIMARY,
        rowheight=26,
        borderwidth=0,
        font=(theme.FONT_FAMILY, 11),
    )
    style.map(
        "Dark.Treeview",
        background=[("selected", theme.COLOR_ACCENT)],
        foreground=[("selected", "#ffffff")],
    )
    style.configure(
        "Dark.Treeview.Heading",
        background=theme.COLOR_BG_SIDEBAR,
        foreground=theme.COLOR_TEXT_PRIMARY,
        font=(theme.FONT_FAMILY, 11, "bold"),
        borderwidth=0,
        relief="flat",
    )
    style.map("Dark.Treeview.Heading", background=[("active", theme.COLOR_ACCENT_SOFT)])


class ScreenHeader(ctk.CTkFrame):
    """رأس موحد لكل شاشة: عنوان + وصف + زر تحديث اختياري."""

    def __init__(self, master, title: str, subtitle: str = "", on_refresh: Callable | None = None, **kwargs):
        super().__init__(master, fg_color="transparent", **kwargs)

        text_frame = ctk.CTkFrame(self, fg_color="transparent")
        text_frame.pack(side="right", fill="y")

        ctk.CTkLabel(
            text_frame, text=title, font=theme.title_font(20), text_color=theme.COLOR_TEXT_PRIMARY,
            anchor="e", justify="right",
        ).pack(anchor="e")

        if subtitle:
            ctk.CTkLabel(
                text_frame, text=subtitle, font=theme.body_font(12), text_color=theme.COLOR_TEXT_SECONDARY,
                anchor="e", justify="right",
            ).pack(anchor="e", pady=(2, 0))

        if on_refresh is not None:
            self.refresh_btn = ctk.CTkButton(
                self, text="⟳ تحديث الآن", command=on_refresh, width=120,
                fg_color=theme.COLOR_ACCENT, hover_color=theme.COLOR_ACCENT_HOVER,
                font=theme.body_font(12, "bold"),
            )
            self.refresh_btn.pack(side="left", anchor="n")


class KPICard(ctk.CTkFrame):
    """بطاقة إحصائية صغيرة (Key Performance Indicator) للوحة التحكم."""

    def __init__(self, master, label: str, value: str = "—", accent: str = theme.COLOR_ACCENT, **kwargs):
        super().__init__(
            master, fg_color=theme.COLOR_BG_PANEL, corner_radius=14,
            border_width=1, border_color=theme.COLOR_BORDER, **kwargs,
        )

        self._bar = ctk.CTkFrame(self, fg_color=accent, width=4, corner_radius=2)
        self._bar.pack(side="right", fill="y", padx=(0, 0), pady=12)

        inner = ctk.CTkFrame(self, fg_color="transparent")
        inner.pack(side="right", fill="both", expand=True, padx=16, pady=14)

        self.value_label = ctk.CTkLabel(
            inner, text=value, font=theme.title_font(26), text_color=theme.COLOR_TEXT_PRIMARY,
            anchor="e", justify="right",
        )
        self.value_label.pack(anchor="e")

        ctk.CTkLabel(
            inner, text=label, font=theme.body_font(12), text_color=theme.COLOR_TEXT_SECONDARY,
            anchor="e", justify="right",
        ).pack(anchor="e", pady=(2, 0))

    def set_value(self, value: str):
        self.value_label.configure(text=value)


class SearchableTable(ctk.CTkFrame):
    """
    جدول بيانات مزود بشريط بحث علوي يفلتر جميع الأعمدة فوراً، وزر تصدير CSV.
    الاستخدام: table.set_columns([...]) ثم table.load_rows([...dicts...])
    """

    def __init__(self, master, columns: Sequence[str] = (), placeholder: str = "ابحث في هذا الجدول...", **kwargs):
        super().__init__(master, fg_color="transparent", **kwargs)
        _style_treeview()

        self._all_rows: List[dict] = []
        self._columns: Sequence[str] = ()

        # --- شريط البحث والأدوات العلوي ---
        toolbar = ctk.CTkFrame(self, fg_color="transparent")
        toolbar.pack(fill="x", pady=(0, 8))

        self.export_btn = ctk.CTkButton(
            toolbar, text="⬇ تصدير Excel", width=125, command=self._export_excel,
            fg_color=theme.COLOR_BG_PANEL_ALT, hover_color=theme.COLOR_ACCENT_SOFT,
            border_width=1, border_color=theme.COLOR_BORDER, font=theme.body_font(12),
        )
        self.export_btn.pack(side="left", padx=(0, 6))

        self.csv_btn = ctk.CTkButton(
            toolbar, text="CSV", width=70, command=self._export_csv,
            fg_color=theme.COLOR_BG_PANEL_ALT, hover_color=theme.COLOR_ACCENT_SOFT,
            border_width=1, border_color=theme.COLOR_BORDER, font=theme.body_font(11),
        )
        self.csv_btn.pack(side="left")

        self.count_label = ctk.CTkLabel(
            toolbar, text="", font=theme.body_font(11), text_color=theme.COLOR_TEXT_MUTED,
        )
        self.count_label.pack(side="left", padx=(10, 0))

        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *_: self._apply_filter())
        search_entry = ctk.CTkEntry(
            toolbar, textvariable=self.search_var, placeholder_text=placeholder,
            width=280, justify="right",
        )
        search_entry.pack(side="right")

        # --- الجدول نفسه ---
        table_frame = ctk.CTkFrame(self, fg_color=theme.COLOR_BG_PANEL, corner_radius=10)
        table_frame.pack(fill="both", expand=True)
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        self.tree = ttk.Treeview(table_frame, style="Dark.Treeview", show="headings", selectmode="browse")
        self.tree.grid(row=0, column=0, sticky="nsew", padx=(1, 0), pady=1)

        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        vsb.grid(row=0, column=1, sticky="ns", pady=1)
        hsb = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        hsb.grid(row=1, column=0, sticky="ew", padx=(1, 0))
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        if columns:
            self.set_columns(columns)

    def set_columns(self, columns: Sequence[str]):
        self._columns = list(columns)
        self.tree.configure(columns=self._columns)
        for col in self._columns:
            self.tree.heading(col, text=col, anchor="e")
            self.tree.column(col, anchor="e", width=140, stretch=True)

    def load_rows(self, rows: List[dict]):
        """rows: قائمة قواميس بمفاتيح تطابق أسماء الأعمدة المضبوطة في set_columns."""
        self._all_rows = rows
        self._apply_filter()

    def clear(self):
        self._all_rows = []
        self._apply_filter()

    def _apply_filter(self):
        query = self.search_var.get().strip().lower()
        for item in self.tree.get_children():
            self.tree.delete(item)

        visible = 0
        for row in self._all_rows:
            values = [str(row.get(col, "")) for col in self._columns]
            if query and not any(query in v.lower() for v in values):
                continue
            self.tree.insert("", "end", values=values)
            visible += 1

        total = len(self._all_rows)
        if query:
            self.count_label.configure(text=f"عرض {visible} من أصل {total}")
        else:
            self.count_label.configure(text=f"الإجمالي: {total}")

    def _get_visible_rows(self):
        query = self.search_var.get().strip().lower()
        if not query:
            return list(self._all_rows)
        visible = []
        for row in self._all_rows:
            values = [str(row.get(col, "")) for col in self._columns]
            if any(query in value.lower() for value in values):
                visible.append(row)
        return visible

    def _export_excel(self):
        rows = self._get_visible_rows()
        if not rows:
            messagebox.showinfo("تصدير Excel", "لا توجد بيانات ظاهرة لتصديرها حالياً.")
            return
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            title="حفظ النتائج كملف Excel منسق",
            initialfile=f"DFIR_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        )
        if not file_path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.worksheet.table import Table, TableStyleInfo
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "DFIR Results"
            ws.sheet_view.rightToLeft = True
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(self._columns))}{len(rows)+1}"

            header_fill = PatternFill("solid", fgColor="1F4E78")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            body_font = Font(name="Segoe UI", size=10, color="000000")
            thin = Side(style="thin", color="D9E1F2")
            border = Border(bottom=thin)

            for col_idx, col in enumerate(self._columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border

            for row_idx, row in enumerate(rows, 2):
                for col_idx, col in enumerate(self._columns, 1):
                    value = row.get(col, "")
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.font = body_font
                    cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                    cell.border = border
                ws.row_dimensions[row_idx].height = 32
            ws.row_dimensions[1].height = 30

            for col_idx, col in enumerate(self._columns, 1):
                values = [str(col)] + [str(r.get(col, "")) for r in rows[:300]]
                max_len = min(max((len(v) for v in values), default=10) + 2, 65)
                ws.column_dimensions[get_column_letter(col_idx)].width = max(12, max_len)

            if rows:
                ref = f"A1:{get_column_letter(len(self._columns))}{len(rows)+1}"
                table = Table(displayName="DFIRResults", ref=ref)
                style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
                table.tableStyleInfo = style
                ws.add_table(table)

            wb.save(file_path)
            messagebox.showinfo("تصدير Excel", f"تم تصدير {len(rows)} صفاً بنجاح إلى:\n{file_path}")
        except Exception as exc:
            messagebox.showerror("خطأ في التصدير", f"تعذر إنشاء ملف Excel:\n{exc}")

    def _export_csv(self):
        rows = self._get_visible_rows()
        if not rows:
            messagebox.showinfo("تصدير CSV", "لا توجد بيانات ظاهرة لتصديرها حالياً.")
            return
        file_path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv")],
            title="حفظ نتائج الجدول كملف CSV",
            initialfile=f"DFIR_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
        )
        if not file_path:
            return
        try:
            with open(file_path, "w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=self._columns)
                writer.writeheader()
                for row in rows:
                    writer.writerow({col: row.get(col, "") for col in self._columns})
            messagebox.showinfo("تصدير CSV", f"تم حفظ {len(rows)} صفاً بنجاح في:\n{file_path}")
        except Exception as exc:
            messagebox.showerror("خطأ في التصدير", f"تعذر حفظ الملف:\n{exc}")


class LoadingBanner(ctk.CTkFrame):
    """شريط صغير يظهر أعلى الشاشة أثناء جلب الأدلة في الخلفية."""

    def __init__(self, master, text: str = "جاري جلب البيانات...", **kwargs):
        super().__init__(master, fg_color=theme.COLOR_ACCENT_SOFT, corner_radius=8, **kwargs)
        self.label = ctk.CTkLabel(
            self, text=text, font=theme.body_font(12, "bold"), text_color=theme.COLOR_INFO,
        )
        self.label.pack(padx=14, pady=8)

    def set_text(self, text: str):
        self.label.configure(text=text)
